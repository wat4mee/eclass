"""Text extraction from downloaded course files (PDF, PPTX, DOCX), page by page."""
import io
import re
from pathlib import Path

from . import db

SUPPORTED = {".pdf": "pdf", ".pptx": "pptx", ".docx": "docx"}
MIN_CHARS_PER_PAGE = 20  # below this on average the file is probably scanned (no text layer)


def _tidy(text):
    lines = [" ".join(line.split()) for line in text.splitlines()]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


OCR_DPI = 200


def _ocr_page(page):
    """OCR one PDF page with the macOS Vision framework (via ocrmac); None if unavailable."""
    try:
        from ocrmac import ocrmac
        from PIL import Image
    except ImportError:
        return None
    import io

    img = Image.open(io.BytesIO(page.get_pixmap(dpi=OCR_DPI).tobytes("png")))
    lines = ocrmac.OCR(img, recognition_level="accurate", language_preference=["en-US"]).recognize()
    # bbox = (x, y, w, h), normalized, origin bottom-left: read top-to-bottom, left-to-right
    lines.sort(key=lambda r: (-round(r[2][1] + r[2][3], 3), r[2][0]))
    return _tidy("\n".join(text for text, _conf, _bbox in lines))


def extract_pdf(path):
    """Text layer per page; image-only pages (screenshots, scans) are OCR'd."""
    import pymupdf

    with pymupdf.open(path) as doc:
        return _pdf_pages(doc)


def _pdf_pages(doc):
    pages, used_ocr = [], False
    for page in doc:
        text = _tidy(page.get_text("text"))
        if len(text) < MIN_CHARS_PER_PAGE and page.get_images():
            ocr = _ocr_page(page)  # None where OCR is unavailable (Linux servers)
            if ocr is not None:
                text, used_ocr = ocr, True
        pages.append(text)
    return pages, ("pdf+ocr" if used_ocr else "pdf")


def _shape_texts(shapes):
    for shape in shapes:
        if shape.shape_type == 6:  # group: recurse
            yield from _shape_texts(shape.shapes)
        elif shape.has_text_frame:
            yield shape.text_frame.text
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                yield " | ".join(cell.text for cell in row.cells)


def extract_pptx(path):
    from pptx import Presentation

    slides = []
    for slide in Presentation(path).slides:
        parts = list(_shape_texts(slide.shapes))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"[Notes] {notes}")
        slides.append(_tidy("\n".join(parts)))
    return slides, "pptx"


def extract_docx(path):
    import docx

    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return [_tidy("\n".join(parts))], "docx"


EXTRACTORS = {"pdf": extract_pdf, "pptx": extract_pptx, "docx": extract_docx}


def extract_bytes(data: bytes, name: str) -> tuple[list[str], str]:
    """Like the extract_* functions, from file contents in memory (the hosted version never stores files)."""
    method = method_for(name)
    if method is None:
        raise ValueError(f"unsupported file type: {Path(name).suffix or name}")
    if method == "pdf":
        import pymupdf

        with pymupdf.open(stream=data, filetype="pdf") as doc:
            return _pdf_pages(doc)
    return EXTRACTORS[method](io.BytesIO(data))  # python-pptx / python-docx read file-like objects


def method_for(name):
    return SUPPORTED.get(Path(name).suffix.lower())


def pending_files(conn, force=False):
    """Supported files whose text is missing, stale, or image-only PDFs not yet OCR'd."""
    where = "" if force else f"""WHERE e.file_id IS NULL OR e.sha256 != f.sha256
        OR (e.method = 'pdf' AND e.error IS NULL
            AND e.n_chars < {MIN_CHARS_PER_PAGE} * e.n_pages)"""
    rows = conn.execute(
        f"""SELECT f.* FROM files f LEFT JOIN extractions e ON e.file_id = f.id
            {where} ORDER BY f.id"""
    ).fetchall()
    return [r for r in rows if method_for(r["path"])]


def extract_pending(conn, force=False, log=print):
    """Extract text for new/changed files. Returns a dict of outcome counts."""
    stats = {"extracted": 0, "failed": 0, "no_text": 0}
    for f in pending_files(conn, force):
        method = method_for(f["path"])
        try:
            (pages, method), error = EXTRACTORS[method](f["path"]), None
        except Exception as exc:  # a corrupt / protected file must not stop the run
            pages, error = [], f"{type(exc).__name__}: {exc}"
        db.save_extraction(conn, f["id"], f["sha256"], method, pages, error)
        chars = sum(len(p) for p in pages)
        if error:
            stats["failed"] += 1
            log(f"  FAIL  {f['filename']}: {error}")
            continue
        stats["extracted"] += 1
        flag = ""
        if chars < MIN_CHARS_PER_PAGE * max(len(pages), 1):
            stats["no_text"] += 1
            flag = "  <- no text layer (scanned?)"
        log(f"  {method:<8}{len(pages):>5} pages{chars:>9} chars  {f['filename']}{flag}")
    return stats
